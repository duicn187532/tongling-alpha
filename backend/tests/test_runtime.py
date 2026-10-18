import io
import time
import zipfile
import pytest
from sqlalchemy import text
from app.db import Session
from app.evidence import parse_file
from app.intelligence import validate_endpoint, resolve_temporal


def run(client, pid):
    response = client.post(f"/api/projects/{pid}/agent/continue")
    assert response.status_code == 200, response.text
    job = response.json()
    for _ in range(200):
        status = client.get("/api/jobs/" + job["id"]).json()
        if status["status"] not in {"RUNNING", "QUEUED"}:
            assert status["status"] in {"SUCCEEDED", "WAITING"}, status
            return client.get("/api/projects/" + pid).json()
        time.sleep(0.05)
    raise AssertionError("Run did not complete")


def wait_meeting(client, mid):
    for _ in range(200):
        meeting = client.get("/api/meetings/" + mid).json()
        if meeting["status"] != "ANALYZING":
            return meeting
        time.sleep(0.02)
    raise AssertionError("Meeting analysis did not complete")


def accept_all(client, project):
    for requirement in project["requirements"]:
        response = client.put(
            "/api/requirements/" + requirement["id"] + "/decision",
            json={"action": "ACCEPT", "rationale": "測試人工確認來源"},
        )
        assert response.status_code == 200, response.text


def test_baseline_revision_download_immutability(client, project):
    p = run(client, project["id"])
    assert p["runtime"] == "WAITING_HUMAN"
    assert client.post(f"/api/projects/{p['id']}/prd/revisions").status_code == 422
    iid = p["interruptions"][0]["id"]
    client.post(f"/api/interruptions/{iid}/dismiss")
    p = client.get("/api/projects/" + p["id"]).json()
    assert p["interruptions"][0]["status"] == "DISMISSED"
    assert p["runtime"] == "WAITING_HUMAN"
    accept_all(client, p)
    draft = client.post(f"/api/projects/{p['id']}/prd/revisions").json()
    assert draft["version"] == "v0.1"
    # A repeated generate request returns the existing working revision.
    assert (
        client.post(f"/api/projects/{p['id']}/prd/revisions").json()["id"]
        == draft["id"]
    )
    response = client.post(f"/api/prd-revisions/{draft['id']}/baseline")
    assert response.status_code == 200, response.text
    original = response.json()
    assert original["version"] == "v1.0"
    download = client.get(f"/api/prd-revisions/{original['id']}/download")
    assert download.status_code == 200
    with zipfile.ZipFile(io.BytesIO(download.content)) as z:
        assert "word/document.xml" in z.namelist()
        assert len([n for n in z.namelist() if n.startswith("word/media/")]) == 3
    with Session() as db:
        with pytest.raises(Exception, match="immutable"):
            db.execute(
                text("UPDATE prd_revisions SET content = '{}' WHERE id = :id"),
                {"id": original["id"]},
            )
            db.commit()
        db.rollback()
        with pytest.raises(Exception, match="immutable"):
            db.execute(
                text("DELETE FROM prd_revisions WHERE id = :id"), {"id": original["id"]}
            )
            db.commit()
    response = client.post(
        f"/api/projects/{p['id']}/requirement-text",
        json={"text": "登入驗證：改為雙因素驗證"},
    )
    assert response.status_code == 200
    updated = run(client, p["id"])
    assert updated["changes"][0]["status"] == "OPEN"
    cid = updated["changes"][0]["id"]
    for mode in ["MANUAL", "AGENT"]:
        client.post(f"/api/projects/{p['id']}/workflow/mode", json={"mode": mode})
    client.post(
        "/api/change-requests/" + cid + "/resolve",
        json={"action": "REVISION_REQUIRED", "rationale": "已確認需要更新驗證方式"},
    )
    for r in updated["requirements"]:
        if r["topic"] == "登入驗證":
            response = client.put(
                "/api/requirements/" + r["id"] + "/decision",
                json={
                    "action": "MODIFY",
                    "statement": "登入驗證：改為雙因素驗證",
                    "rationale": "正式決策",
                },
            )
            assert response.status_code == 200
    response = client.post(f"/api/projects/{p['id']}/prd/revisions")
    assert response.status_code == 200, response.text
    next_revision = response.json()
    assert next_revision["version"] == "v1.1"
    assert next_revision["parent_revision_id"] == original["id"]
    assert len(next_revision["diff"]) == 1
    assert next_revision["content"]["background"] == original["content"]["background"]
    unchanged = lambda r: next(
        x for x in r["content"]["requirements"] if x["topic"] == "查詢紀錄"
    )
    assert unchanged(next_revision) == unchanged(original)
    response = client.post(f"/api/prd-revisions/{next_revision['id']}/baseline")
    assert response.status_code == 200, response.text
    old_download = client.get(f"/api/prd-revisions/{original['id']}/download")
    assert old_download.content == download.content
    final = client.get("/api/projects/" + p["id"]).json()
    old = next(r for r in final["revisions"] if r["id"] == original["id"])
    assert old["content_hash"] == original["content_hash"]
    assert old["status"] == "SUPERSEDED"
    assert (
        client.get(f"/api/prd-revisions/{next_revision['id']}/download").status_code
        == 200
    )


def test_local_only_and_settings(client):
    with pytest.raises(ValueError, match="Local-only"):
        validate_endpoint(
            {
                "provider": "openai_compatible",
                "endpoint": "https://example.org/v1",
                "model": "test",
                "local_only": True,
            }
        )
    validate_endpoint(
        {
            "provider": "ollama",
            "endpoint": "http://127.0.0.1:11434",
            "model": "test",
            "local_only": True,
        }
    )
    assert (
        client.put(
            "/api/settings",
            json={
                "provider": "openai_compatible",
                "model": "test",
                "endpoint": "https://example.org/v1",
                "local_only": True,
            },
        ).status_code
        == 422
    )


def test_upload_parse_search_attach(client, project):
    response = client.post(
        "/api/knowledge/upload",
        files={
            "file": (
                "policy.md",
                "驗證規範：所有交易使用一次性驗證碼".encode(),
                "text/markdown",
            )
        },
    )
    assert response.status_code == 200, response.text
    source = response.json()
    results = client.post(
        f"/api/projects/{project['id']}/knowledge/search", json={"query": "交易 驗證碼"}
    ).json()
    assert source["id"] in [r["source_id"] for r in results]
    p = client.get("/api/projects/" + project["id"]).json()
    assert source["id"] not in [s["id"] for s in p["sources"]]
    client.post(
        f"/api/projects/{project['id']}/knowledge/attach",
        json={"source_id": source["id"]},
    )
    assert source["id"] in [
        s["id"] for s in client.get("/api/projects/" + project["id"]).json()["sources"]
    ]
    assert (
        client.post(
            "/api/knowledge/upload",
            files={"file": ("bad.exe", b"MZ", "application/octet-stream")},
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/knowledge/upload",
            files={"file": ("bad.pdf", b"not a PDF", "application/pdf")},
        ).status_code
        == 422
    )


def test_folder_incremental_and_removed(client, project, tmp_path):
    (tmp_path / "rule.md").write_text("驗證：需要驗證碼", encoding="utf8")
    connector = client.post(
        "/api/connectors", json={"path": str(tmp_path), "project_id": project["id"]}
    ).json()
    endpoint = "/api/connectors/" + connector["id"] + "/scan"
    assert client.post(endpoint).json()["changed"] == 1
    assert client.post(endpoint).json()["changed"] == 0
    (tmp_path / "rule.md").write_text("驗證：需要雙因素驗證", encoding="utf8")
    assert client.post(endpoint).json()["changed"] == 1
    (tmp_path / "rule.md").unlink()
    assert client.post(endpoint).json()["removed"] == 1
    sources = client.get("/api/projects/" + project["id"]).json()["sources"]
    assert {"SUPERSEDED", "REMOVED"}.issubset({s["status"] for s in sources})


def test_temporal_reopen_silence():
    def fact(day, kind, statement):
        return {"meeting_date": day, "kind": kind, "statement": statement}

    current, reopened, history = resolve_temporal(
        [
            fact("2026-10-15", "CHANGE", "新決議"),
            fact("2026-10-01", "DECISION", "舊決議"),
            fact("2026-10-22", "OPEN_ITEM", "待釐清"),
        ]
    )
    assert {x["statement"] for x in current} == {"舊決議", "新決議"} and reopened
    assert history[0]["validity"] == "CURRENT"
    current, reopened, _ = resolve_temporal(
        [
            fact("2026-10-01", "DECISION", "原決議"),
            fact("2026-10-22", "QUESTION", "提問"),
        ]
    )
    assert current[0]["statement"] == "原決議" and not reopened


def test_meetings_wait_resume(client, project):
    pid = project["id"]
    assert (
        client.post(
            f"/api/projects/{pid}/meetings", json={"title": "缺少日期"}
        ).status_code
        == 422
    )
    client.post(f"/api/projects/{pid}/wait-meeting")
    assert run(client, pid)["runtime"] == "WAITING_MEETING"
    for day in ["2026-10-15", "2026-10-01", "2026-10-08"]:
        response = client.post(
            f"/api/projects/{pid}/meetings",
            json={
                "title": "需求會議",
                "meeting_date": day,
                "transcript": "驗證方法：決議使用驗證碼",
            },
        )
        assert response.status_code == 200, response.text
        meeting = response.json()
        assert (
            client.post("/api/meetings/" + meeting["id"] + "/analyze").status_code
            == 200
        )
        assert (
            wait_meeting(client, meeting["id"])
            and client.post("/api/meetings/" + meeting["id"] + "/confirm")
        ).status_code == 200
    p = run(client, pid)
    assert [m["meeting_date"] for m in p["meetings"]] == [
        "2026-10-01",
        "2026-10-08",
        "2026-10-15",
    ]
    assert p["state"]["waiting_meeting"] is False
    assert p["runtime"] == "WAITING_HUMAN"


def test_parsers_all_formats():
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    import fitz

    doc = Document()
    doc.add_paragraph("Requirement source")
    buf = io.BytesIO()
    doc.save(buf)
    assert "Requirement" in parse_file("x.docx", buf.getvalue())[0]["text"]
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide.shapes.title.text = "Requirement source"
    buf = io.BytesIO()
    ppt.save(buf)
    assert "Requirement" in parse_file("x.pptx", buf.getvalue())[0]["text"]
    wb = Workbook()
    wb.active.append(["Requirement", "source"])
    buf = io.BytesIO()
    wb.save(buf)
    assert "Requirement" in parse_file("x.xlsx", buf.getvalue())[0]["text"]
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Requirement source")
    assert "Requirement" in parse_file("x.pdf", pdf.tobytes())[0]["text"]
    for ext in ["txt", "md", "csv"]:
        assert parse_file("x." + ext, b"Requirement, source")
    assert parse_file(
        "x.vtt", b"WEBVTT\n\n00:00:00.000 --> 00:00:01.000\nSpeaker: decision\n"
    )[0]["time"]


def test_epoch_cancels_stale_worker(client, project, monkeypatch):
    import threading
    import app.workflow as workflow

    entered, release = threading.Event(), threading.Event()
    real = workflow.extract_sources

    def slow(*args):
        entered.set()
        release.wait(5)
        return real(*args)

    monkeypatch.setattr(workflow, "extract_sources", slow)
    job = client.post(f"/api/projects/{project['id']}/agent/continue").json()
    assert entered.wait(3)
    client.post(f"/api/projects/{project['id']}/agent/pause")
    release.set()
    for _ in range(100):
        status = client.get("/api/jobs/" + job["id"]).json()
        if status["status"] != "RUNNING":
            break
        time.sleep(0.03)
    assert status["status"] == "CANCELLED"
    p = client.get("/api/projects/" + project["id"]).json()
    assert p["runtime"] == "PAUSED" and not p["requirements"]


def test_host_rejects_cross_origin_writes_and_pairing(client):
    # The host process is loopback-only; remote access goes through the gateway.
    assert client.post("/api/pair", json={"code": "wrong"}).status_code == 404
    assert (
        client.post(
            "/api/projects",
            json={"name": "blocked"},
            headers={"origin": "https://malicious.example"},
        ).status_code
        == 403
    )
    assert client.get("/api/projects", headers={"host": "192.168.1.5:8000"}).status_code == 400


def test_no_change_keeps_baseline_and_source_trace(
    client, project, monkeypatch, tmp_path
):
    import app.lifecycle as lifecycle

    # The full renderer is covered by the lifecycle integration test above.
    output = tmp_path / "test.docx"
    output.write_bytes(b"test artifact")
    monkeypatch.setattr(lifecycle, "render_docx", lambda *_: str(output))
    p = run(client, project["id"])
    accept_all(client, p)
    draft = client.post(f"/api/projects/{p['id']}/prd/revisions").json()
    original = client.post(f"/api/prd-revisions/{draft['id']}/baseline").json()
    client.post(
        f"/api/projects/{p['id']}/requirement-text",
        json={"text": "登入驗證：不採用的新建議"},
    )
    updated = run(client, p["id"])
    cid = updated["changes"][0]["id"]
    response = client.post(
        "/api/change-requests/" + cid + "/resolve",
        json={"action": "NO_PRD_CHANGE", "rationale": "此建議不適用"},
    )
    assert response.status_code == 200, response.text
    final = run(client, p["id"])
    assert final["runtime"] == "BASELINED"
    assert len(final["revisions"]) == 1
    assert final["revisions"][0]["content_hash"] == original["content_hash"]
    assert final["changes"][0]["status"] == "NO_PRD_CHANGE"
    assert len(final["sources"]) == len(p["sources"]) + 1


def test_stale_draft_cannot_be_baselined(client, project):
    p = run(client, project["id"])
    accept_all(client, p)
    draft = client.post(f"/api/projects/{p['id']}/prd/revisions").json()
    client.post(
        f"/api/projects/{p['id']}/requirement-text",
        json={"text": "新證據：需要重新審查"},
    )
    response = client.post(f"/api/prd-revisions/{draft['id']}/baseline")
    assert response.status_code == 422
