"""Slides give up their tables, notes and pictures; images need a vision model."""

import io

from app import vision
from app.evidence import LocalFolderProvider, parse_file

PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 64


def deck_with_table_and_notes():
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "提案效益"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "指標", "目標"
    table.cell(1, 0).text, table.cell(1, 1).text = "客訴件數", "下降 30%"
    slide.notes_slide.notes_text_frame.text = "數字來自 2026 年客訴統計"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_slide_tables_and_speaker_notes_are_read():
    text = parse_file("proposal.pptx", deck_with_table_and_notes())[0]["text"]
    assert "提案效益" in text and "客訴件數 | 下降 30%" in text and "備註：數字來自 2026 年客訴統計" in text


def test_image_upload_without_vision_model_explains_why(client, project):
    response = client.post(f"/api/projects/{project['id']}/sources/upload",
                           files={"file": ("wireframe.png", PNG, "image/png")})
    assert response.status_code == 422 and "不支援圖片判讀" in response.json()["detail"]


def test_image_upload_becomes_a_labelled_source(client, project, monkeypatch):
    monkeypatch.setattr(vision, "pick_model", lambda config: {"provider": "ollama", "endpoint": "x", "model": "gemma4:12b"})
    monkeypatch.setattr(vision, "read_image", lambda config, data, selected=None: vision.ScreenReading(screens=[
        vision.Screen(name="薪資通知設定", purpose="員工設定附件密碼", fields=["密碼規則", "手機號碼"], actions=["儲存"])]))
    response = client.post(f"/api/projects/{project['id']}/sources/upload",
                           files={"file": ("wireframe.png", PNG, "image/png")}, data={"kind": "USER_UPLOAD"})
    assert response.status_code == 200, response.text
    content = response.json()["content"]
    assert content.startswith("【wireframe.png・AI 圖片判讀（gemma4:12b），內容待確認】")
    assert "欄位：密碼規則、手機號碼" in content and "操作：儲存" in content


def test_mismatched_image_bytes_are_rejected():
    try:
        parse_file("fake.png", b"not an image")
    except ValueError as exc:
        assert "圖片內容與副檔名不符" in str(exc)
    else:
        raise AssertionError("expected rejection")


def test_folder_scan_leaves_images_for_upload(tmp_path):
    (tmp_path / "spec.md").write_text("規格", encoding="utf-8")
    (tmp_path / "screen.png").write_bytes(PNG)
    assert LocalFolderProvider(tmp_path).list_objects() == ["spec.md"]
