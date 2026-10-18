"""The standalone company simulator indexes its own folder through an API."""

from fastapi.testclient import TestClient
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from company_file_server import app as service


def test_company_file_api_search_rescan_and_removal(tmp_path, monkeypatch):
    folder = tmp_path / "company"
    folder.mkdir()
    document = folder / "invoice-policy.md"
    document.write_text("電子帳單申請需要通知信箱驗證與受理編號。", encoding="utf-8")
    monkeypatch.setattr(service, "index", service.FileIndex(folder))
    with TestClient(service.app) as client:
        assert client.get("/health").json()["documents"] == 1
        assert "SEARCH" in client.get("/capabilities").json()["capabilities"]
        hit = client.post("/search", json={"query": "電子帳單", "limit": 3}).json()["results"][0]
        assert hit["title"] == document.name
        assert hit["metadata"]["verified_authority"] is False
        evidence = client.get(f"/evidence/{hit['id']}").json()
        assert "受理編號" in evidence["content"]
        assert client.get("/documents").json()["documents"][0]["id"] == hit["id"]
        document.unlink()
        assert client.post("/rescan").json()["documents"] == 0
        assert client.post("/search", json={"query": "電子帳單", "limit": 3}).json()["results"] == []


def test_company_file_server_rejects_other_host_and_bad_file(tmp_path, monkeypatch):
    folder = tmp_path / "company"
    folder.mkdir()
    (folder / "bad.pdf").write_bytes(b"not a PDF")
    monkeypatch.setattr(service, "index", service.FileIndex(folder))
    with TestClient(service.app) as client:
        health = client.get("/health").json()
        assert health["documents"] == 0 and health["errors"] == 1
        assert client.get("/documents", headers={"Host": "untrusted.example"}).status_code == 400


def test_company_file_server_graph_image_and_office(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    from openpyxl import Workbook

    folder = tmp_path / "company"
    folder.mkdir()
    (folder / "流程.graphml").write_text('<graphml><graph><node id="審核節點"/><edge source="申請" target="審核"/></graph></graphml>', encoding="utf-8")
    Image.new("RGB", (24, 16), "#ab4567").save(folder / "流程圖.png")
    doc = Document()
    doc.add_paragraph("電子帳單申請作業")
    doc.save(folder / "作業說明.docx")
    workbook = Workbook()
    workbook.active.append(["通知信箱", "一次性驗證碼"])
    workbook.save(folder / "欄位清單.xlsx")
    (folder / "附件.dat").write_bytes(b"\x00\x01\x02")
    monkeypatch.setattr(service, "index", service.FileIndex(folder))
    with TestClient(service.app) as client:
        rows = client.get("/documents").json()["documents"]
        assert len(rows) == 5
        graph = next(row for row in rows if row["category"] == "GRAPH")
        image = next(row for row in rows if row["category"] == "IMAGE")
        assert image["content_mode"] == "METADATA_ONLY"
        assert next(row for row in rows if row["file_type"] == "dat")["category"] == "FILE"
        assert client.post("/graph/search", json={"query": "審核", "limit": 3}).json()["results"][0]["id"] == graph["id"]
        assert client.post("/search", json={"query": "驗證碼", "limit": 3}).json()["results"]
        assert client.get(f"/files/{graph['id']}").status_code == 200
        preview = client.get(f"/preview/{image['id']}")
        assert preview.status_code == 200 and preview.headers["content-type"].startswith("image/png")
        assert client.get(f"/preview/{graph['id']}").status_code == 404
