import io
import numpy as np
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from app.db import Base, Workspace, Project
from app.evidence import ingest
from app.retrieval_v2 import (
    RetrievalPipelineV2,
    normalize,
    aggregate,
    analyze_query,
    metadata_adjustment,
    matches_filters,
    rerank_passage,
    missing_exact_identifier,
)
from app.retrieval_structure import (
    parse_structured,
    structural_chunks,
    logical_document,
)
from app.retrieval_providers import local_endpoint, Cache, LocalCrossEncoderReranker


@pytest.fixture
def corpus():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        db.add(Workspace(id="w", name="w"))
        db.flush()
        for pid in ["a", "b"]:
            db.add(Project(id=pid, workspace_id="w", name=pid))
        db.flush()
        for i in range(6):
            ingest(
                db,
                f"付款規格{i}",
                [{"text": f"付款 HTTP 409 API_PAYMENT_01 規則 {i}"}] * 8,
                project=db.get(Project, "a"),
                meta={"source_path": f"{i}.md", "project_id": "a"},
            )
        ingest(
            db,
            "隔離文件",
            [{"text": "付款 HTTP 409 API_PAYMENT_01"}],
            project=db.get(Project, "b"),
        )
        yield db
    engine.dispose()


def test_candidate_expansion_and_scope(corpus, tmp_path):
    pipe = RetrievalPipelineV2(tmp_path)
    output = pipe.search(
        corpus,
        "付款 HTTP 409",
        "a",
        {"top_k": 6, "retrieval_v2": {"lexical_chunks": 2, "candidate_documents": 6}},
    )
    assert len(output["results"]) == 6
    assert output["trace"]["actual_chunk_count"] > 2
    assert all(r["title"] != "隔離文件" for r in output["results"])
    assert len({r["source_id"] for r in output["results"]}) == 6
    assert all(r["semantic_score"] is None for r in output["results"])


@pytest.mark.parametrize("method", ["minmax", "rank", "zscore", "rrf"])
def test_normalization(method):
    assert np.isfinite(normalize([1, 1, 1], method)).all()
    assert list(normalize([1, 2, 3], method)) == sorted(normalize([1, 2, 3], method))
    with pytest.raises(ValueError):
        normalize([1, np.nan], method)


def test_aggregation():
    assert aggregate([0.9, 0.8, 0.7], "max") == 0.9
    assert aggregate([0.9, 0.8, 0.7], "top_n_mean") == pytest.approx(0.8)
    assert aggregate([0.9, 0.8, 0.7], "weighted_top_n") == pytest.approx(0.85)


def test_cross_encoder_raw_logits_are_transformed_consistently(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace
    import app.retrieval_providers as providers

    class FakeIdentity:
        pass

    monkeypatch.setitem(sys.modules, "torch", SimpleNamespace(
        nn=SimpleNamespace(Identity=FakeIdentity)))

    model_path = tmp_path / "model"
    model_path.mkdir()
    (model_path / "config.json").write_text("{}", encoding="utf-8")

    class FakeModel:
        def predict(self, pairs, **kwargs):
            assert len(pairs) == 3
            assert isinstance(kwargs["activation_fn"], FakeIdentity)
            return np.array([-0.1, 0.2, 1.2])

    monkeypatch.setattr(providers, "_local_cross_encoder", lambda *args: FakeModel())
    reranker = LocalCrossEncoderReranker(
        {"model_path": str(model_path), "revision": "test-revision"},
        tmp_path / "cache.db",
    )
    scores = reranker.rerank(
        "query",
        [{"id": str(i), "title": str(i), "text": str(i)} for i in range(3)],
    )
    assert list(scores.values()) == pytest.approx(
        [1 / (1 + np.exp(-x)) for x in [-0.1, 0.2, 1.2]]
    )
    assert scores["0"] < scores["1"] < scores["2"]


def test_reranker_passage_modes_use_requested_context():
    doc = {"title": "交易規格", "text": "目前交易段落", "meta": {"heading": "手續費"},
           "evidence_chunks": [{"text": "目前交易段落"}],
           "_parent_context": "前段\n目前交易段落\n後段"}
    assert rerank_passage(doc, "chunk") == "目前交易段落"
    context = rerank_passage(doc, "context")
    assert "交易規格" in context and "手續費" in context
    assert "前段" in context and "後段" in context


def test_exact_identifier_admission_guard():
    doc = {"title":"介面規格", "evidence_chunks":[{"text":"交易欄位 TX_CODE 長度 8"}]}
    assert not missing_exact_identifier("TX_CODE 長度？", doc)
    assert missing_exact_identifier("TLX_UNLISTED_830001 長度？", doc)
    assert not missing_exact_identifier("TX_CODE 和 TRACE_NO 差異？", doc)
    assert not missing_exact_identifier("API-to-host mapping？", doc)
    assert missing_exact_identifier("CR-001 的規格？", doc)


def test_query_temporal_authority():
    q = analyze_query("目前 MAX_BATCH_ITEMS baseline")
    assert {"technical", "exact", "current"} <= set(q["labels"])
    current = metadata_adjustment(
        {"lifecycle_status": "CURRENT", "authority": "high"}, q, 0.6, True, True
    )
    old = metadata_adjustment(
        {"lifecycle_status": "SUPERSEDED", "authority": "high"}, q, 0.6, True, True
    )
    assert current[1] > old[1]
    assert metadata_adjustment({"authority": "high"}, q, 0.1, True, True) == (0, 0)
    assert (
        metadata_adjustment(
            {"lifecycle_status": "SUPERSEDED"}, analyze_query("舊版"), 0.6, True, True
        )[1]
        > 0
    )
    assert (
        metadata_adjustment(
            {"lifecycle_status": "SUPERSEDED"},
            analyze_query("舊版與目前差異"),
            0.6,
            True,
            True,
        )[1]
        == 0
    )
    assert metadata_adjustment(
        {"lifecycle_status": "CURRENT", "authority": "high"},
        analyze_query("交易規則有哪些"), 0.6, True, True,
    ) == (0.04, 0.0)
    assert metadata_adjustment(
        {"authority": "high"}, q, 0.4, True, False,
    ) == (0.0, 0.0)


def test_structure_and_stable_ids():
    from docx import Document

    doc = Document()
    doc.add_heading("Limits", 1)
    doc.add_paragraph("MAX_BATCH_ITEMS HTTP 409")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "field"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "account_id"
    table.cell(1, 1).text = "TEST"
    stream = io.BytesIO()
    doc.save(stream)
    blocks = parse_structured("spec.docx", stream.getvalue())
    chunks = structural_chunks(blocks, "spec", "stable-doc")
    assert any(c["meta"]["heading_path"] == ["Limits"] for c in chunks)
    row = next(c for c in chunks if c["meta"].get("kind") == "table_row")
    assert (
        "field | value" in row["text"]
        and "account_id" in row["meta"]["technical_tokens"]
    )
    assert chunks == structural_chunks(blocks, "spec", "stable-doc")


def test_local_only_cache(tmp_path):
    with pytest.raises(ValueError):
        local_endpoint("https://example.com")
    with pytest.raises(ValueError):
        local_endpoint("http://user:secret@localhost")
    assert local_endpoint("http://127.0.0.1:11434")
    path = tmp_path / "cache.db"
    Cache(path).put_many({"x": [1, 2]})
    assert Cache(path).get_many(["x"]) == {"x": [1, 2]}


def test_reranker_success_and_fallback(corpus, tmp_path):
    class Ranker:
        def rerank(self, q, docs):
            return {d["id"]: i / len(docs) for i, d in enumerate(docs)}

        def metadata(self):
            return {"provider": "test-double"}

    config = {"top_k": 6, "retrieval_v2": {"rerank": True, "candidate_documents": 6}}
    out = RetrievalPipelineV2(tmp_path, reranker_provider=Ranker()).search(
        corpus, "付款", "a", config
    )
    assert out["trace"]["rerank_applied"]
    assert out["results"][0]["rerank_score"] > out["results"][-1]["rerank_score"]

    class Bad:
        def rerank(self, *args):
            raise RuntimeError("unavailable")

    out = RetrievalPipelineV2(tmp_path, reranker_provider=Bad()).search(
        corpus, "付款", "a", config
    )
    assert not out["trace"]["rerank_applied"]
    assert out["trace"]["fallbacks"] and all(
        r["rerank_score"] is None for r in out["results"]
    )


def test_logical_identity_independent_uuid(corpus):
    from types import SimpleNamespace

    a = SimpleNamespace(
        meta={"source_path": "a/spec.md"}, provider="LOCAL", title="X", id="random1"
    )
    b = SimpleNamespace(
        meta={"source_path": "a/spec.md"}, provider="LOCAL", title="X", id="random2"
    )
    assert logical_document(a, "a") == logical_document(b, "a")
    assert logical_document(a, "a") != logical_document(b, "b")
    assert matches_filters(a, {"version": "1"}) is False


def test_superseded_searchable_removed_not(corpus, tmp_path):
    from sqlalchemy import select
    from app.db import Source

    rows = list(corpus.scalars(select(Source)))
    rows[0].status = "SUPERSEDED"
    rows[1].status = "REMOVED"
    corpus.flush()
    config = {"top_k": 20, "retrieval_v2": {"temporal": True}}
    out = RetrievalPipelineV2(tmp_path).search(corpus, "舊版付款", "a", config)
    assert rows[0].id in {r["source_id"] for r in out["results"]}
    assert rows[1].id not in {r["source_id"] for r in out["results"]}


def test_embedding_provider_cache_and_model_version(tmp_path):
    from app.retrieval_providers import OllamaEmbedding

    class FakeLocal(OllamaEmbedding):
        calls = 0

        def request(self, method, path, **kwargs):
            if path == "/api/tags":
                return {"models": [{"name": "local", "digest": self.config["digest"]}]}
            self.calls += 1
            assert kwargs["json"]["truncate"] is False
            return {"embeddings": [[1.0, 2.0] for _ in kwargs["json"]["input"]]}

    cfg = {"model": "local", "digest": "v1"}
    path = tmp_path / "e.db"
    first = FakeLocal(cfg, path)
    first.embed_documents(["same"])
    assert first.calls == 1
    second = FakeLocal(cfg, path)
    second.embed_documents(["same"])
    assert second.calls == 0
    second.embed_query("same")
    assert second.calls == 1
    third = FakeLocal({**cfg, "digest": "v2"}, path)
    third.embed_documents(["same"])
    assert third.calls == 1


def test_debug_api_and_v1_compatibility(client, project):
    response = client.post(
        "/api/projects/" + project["id"] + "/knowledge/debug-search",
        json={"query": "驗證"},
    )
    assert response.status_code == 200, response.text
    assert "trace" in response.json()
    regular = client.post(
        "/api/projects/" + project["id"] + "/knowledge/search", json={"query": "驗證"}
    )
    assert regular.status_code == 200 and isinstance(regular.json(), list)
