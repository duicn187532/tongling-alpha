"""Opt-in structural parsing. Old source/chunk rows remain immutable."""

import io
import json
import re
from pathlib import Path
from .retrieval_providers import fingerprint


def logical_document(source, project_id=None):
    meta = source.meta
    key = (
        meta.get("external_id")
        or meta.get("source_path")
        or meta.get("path")
        or meta.get("relative_path")
        or meta.get("filename")
        or source.title
    )
    # Never use benchmark_doc_id or query labels as identity/ranking features.
    return fingerprint(
        {
            "provider": source.provider,
            "scope": meta.get("project_id", project_id),
            "connector": meta.get("connector_root", meta.get("connector_id", "")),
            "key": str(key).replace("\\", "/").casefold(),
        }
    )


def parse_structured(name, data):
    from .evidence import parse_file

    ext = Path(name).suffix.lower()
    if ext == ".docx":
        # Reuse existing signature/size validation before structural parsing.
        parse_file(name, data)
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        doc = Document(io.BytesIO(data))
        blocks = []
        headings = []
        pidx = 0
        tidx = 0
        for child in doc.element.body.iterchildren():
            if child.tag.endswith("}p"):
                p = Paragraph(child, doc)
                pidx += 1
                if not p.text.strip():
                    continue
                match = re.search(r"Heading (\d+)", p.style.name or "")
                if match:
                    level = int(match[1])
                    headings = headings[: level - 1] + [p.text]
                blocks.append(
                    {
                        "text": p.text,
                        "heading_path": list(headings),
                        "paragraph_index": pidx,
                        "kind": "heading" if match else "paragraph",
                    }
                )
            elif child.tag.endswith("}tbl"):
                table = Table(child, doc)
                tidx += 1
                rows = [[c.text for c in r.cells] for r in table.rows]
                if not rows:
                    continue
                for n, row in enumerate(rows[1:] or rows):
                    blocks.append(
                        {
                            "text": " | ".join(row),
                            "kind": "table_row",
                            "table_name": f"Table {tidx}",
                            "column_headers": rows[0],
                            "row_headers": row[:1],
                            "row_index": n,
                            "heading_path": list(headings),
                        }
                    )
        return blocks
    if ext in {".xlsx", ".csv"}:
        parse_file(name, data)
        if ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = [(s.title, list(s.iter_rows(values_only=True))) for s in wb]
            wb.close()
        else:
            import csv

            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp950")
            sheets = [(Path(name).stem, list(csv.reader(io.StringIO(text))))]
        blocks = []
        for title, rows in sheets:
            rows = [[str(v) if v is not None else "" for v in r] for r in rows]
            rows = [r for r in rows if any(v.strip() for v in r)]
            if not rows:
                continue
            for i, row in enumerate(rows[1:] or rows):
                blocks.append(
                    {
                        "text": " | ".join(row),
                        "kind": "table_row",
                        "table_name": title,
                        "column_headers": rows[0],
                        "row_headers": row[:1],
                        "row_index": i,
                        "heading_path": [title],
                    }
                )
        return blocks
    return parse_file(name, data)


def structural_chunks(blocks, title, document_id, max_chars=1000):
    output = []
    pending = []
    context = None
    occurrences = {}

    def add(text, meta):
        if not text.strip():
            return
        digest = fingerprint({"heading": meta.get("heading_path", []), "text": text})
        n = occurrences.get(digest, 0)
        occurrences[digest] = n + 1
        output.append(
            {
                "text": text,
                "logical_chunk_id": f"{document_id}:{digest}:{n}",
                "meta": {
                    **meta,
                    "document_title": title,
                    "chunk_index": len(output),
                    "content_hash": fingerprint(text),
                    "technical_tokens": re.findall(
                        r"\b[A-Za-z][A-Za-z0-9]*(?:[_-][A-Za-z0-9]+)+\b|HTTP\s+\d{3}",
                        text,
                    ),
                },
            }
        )

    def flush():
        if pending:
            add("\n".join(pending), context or {})
            pending.clear()

    for block in blocks:
        meta = {k: v for k, v in block.items() if k != "text"}
        meta.setdefault(
            "heading_path", [meta["heading"]] if meta.get("heading") else []
        )
        text = block["text"].strip()
        if meta.get("kind") == "table_row":
            flush()
            headers = meta.get("column_headers", [])
            # Preserve each row whole, including repeated schema headers.
            add("Columns: " + " | ".join(headers) + "\n" + text, meta)
            context = None
            continue
        if context != meta or sum(map(len, pending)) + len(text) > max_chars:
            flush()
            context = meta
        if len(text) > max_chars:
            flush()
            # Split at whitespace/punctuation boundaries, never inside a technical token.
            pieces = re.findall(r"[^\s。；;]+[\s。；;]*", text)
            value = ""
            for piece in pieces:
                if value and len(value) + len(piece) > max_chars:
                    add(value, meta)
                    value = ""
                value += piece
            add(value, meta)
        else:
            pending.append(text)
    flush()
    return output


def sidecar_path(root, source):
    return Path(root) / "chunks" / f"{source.id}-{source.sha256}.json"


def write_sidecar(root, source, blocks, project_id=None):
    path = sidecar_path(root, source)
    path.parent.mkdir(parents=True, exist_ok=True)
    chunks = structural_chunks(
        blocks, source.title, logical_document(source, project_id)
    )
    payload = {"schema": 2, "source_hash": source.sha256, "chunks": chunks}
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    temp.replace(path)
    return chunks
