"""Read screens out of wireframes, mockups and slide pictures (AI-11).

The reading becomes ordinary source text, labelled as an AI reading to be
confirmed, so it flows through extraction, evidence and the faithfulness
check like any other document. Only a model that advertises the `vision`
capability is used; otherwise callers get None and say so.
"""

import base64
import io
import json

import httpx
from pydantic import BaseModel, Field

IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp"}
MAX_SLIDE_PICTURES = 4
MIN_PICTURE_BYTES = 20_000  # skip logos and icons

INSTRUCTION = (
    "你是金融業 PM 的助理，正在閱讀一張系統畫面、wireframe 或簡報圖片。"
    "只描述圖片中看得到的內容：每個畫面的名稱與用途、欄位或顯示的資料、使用者可做的操作。"
    "看不清楚的字寫「無法辨識」，不要猜測，不要補上圖片沒有的欄位、規則或數值。"
    "欄位只寫名稱與畫面上的提示文字；勾選框、開關等狀態除非清楚可見，否則不要描述。"
    "若圖片不是系統畫面（例如照片、圖表），在 notes 簡述內容即可。輸出繁體中文。"
)


class Screen(BaseModel):
    name: str = Field(max_length=80)
    purpose: str = Field(default="", max_length=300)
    fields: list[str] = Field(default_factory=list, max_length=40)
    actions: list[str] = Field(default_factory=list, max_length=20)


class ScreenReading(BaseModel):
    screens: list[Screen] = Field(default_factory=list, max_length=6)
    notes: str = Field(default="", max_length=600)


def is_image(data):
    return data.startswith(b"\x89PNG") or data.startswith(b"\xff\xd8") or (data[:4] == b"RIFF" and data[8:12] == b"WEBP")


def pick_model(config):
    """First configured Ollama model that advertises vision, or None."""
    from .model_runtime import model_info, route
    candidates = [route(config, "PRD_ASSIST"), route(config, "SOURCE_FACT_EXTRACTION"),
                  {"provider": config.get("provider"), "endpoint": config.get("endpoint"), "model": config.get("model")}]
    seen = set()
    for selected in candidates:
        key = (selected.get("provider"), selected.get("endpoint"), selected.get("model"))
        if key in seen or selected.get("provider") != "ollama" or not selected.get("model"):
            continue
        seen.add(key)
        try:
            if "vision" in model_info(selected)["capabilities"]:
                return selected
        except Exception:
            continue
    return None


def read_image(config, data, selected=None):
    """A ScreenReading for one image, or None when no vision model is available."""
    from .model_runtime import acquire, release
    selected = selected or pick_model(config)
    if not selected:
        return None
    acquire(config, heavy=False, timeout_seconds=180)
    try:
        with httpx.Client(timeout=httpx.Timeout(240, connect=5), trust_env=False, follow_redirects=False) as client:
            response = client.post(selected["endpoint"].rstrip("/") + "/api/chat", json={
                "model": selected["model"], "stream": False, "format": ScreenReading.model_json_schema(),
                "options": {"temperature": 0, "seed": 42, "num_predict": 1500},
                "messages": [{"role": "system", "content": INSTRUCTION},
                             {"role": "user", "content": "請閱讀這張圖片並依 schema 回傳 JSON。",
                              "images": [base64.b64encode(data).decode()]}]})
            response.raise_for_status()
        return ScreenReading.model_validate(json.loads(response.json()["message"]["content"]))
    finally:
        release()


def reading_text(label, reading, model):
    lines = [f"【{label}・AI 圖片判讀（{model}），內容待確認】"]
    for screen in reading.screens:
        lines.append(f"畫面：{screen.name}" + (f"（{screen.purpose}）" if screen.purpose else ""))
        if screen.fields:
            lines.append("欄位：" + "、".join(screen.fields))
        if screen.actions:
            lines.append("操作：" + "、".join(screen.actions))
    if reading.notes:
        lines.append("備註：" + reading.notes)
    if len(lines) == 1:
        lines.append("圖片中沒有可辨識的畫面或文字內容。")
    return "\n".join(lines)


def slide_pictures(data):
    """(slide number, image bytes) for sizeable pictures in a .pptx."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    found = []
    for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and len(shape.image.blob) >= MIN_PICTURE_BYTES:
                found.append((number, shape.image.blob))
    return found[:MAX_SLIDE_PICTURES]


def image_blocks(name, data, config):
    """Text blocks for an uploaded image or a deck's pictures.

    Images need a vision model and raise when none is available; a deck's
    pictures are optional, so a deck without one just gets a note.
    """
    from pathlib import Path
    ext = Path(name).suffix.lower()
    if ext not in IMAGE_TYPES and ext != ".pptx":
        return []
    selected = pick_model(config)
    if ext in IMAGE_TYPES:
        if not selected:
            raise ValueError("目前設定的模型不支援圖片判讀；請在主機設定選用具影像能力的本地模型（如 Gemma、Qwen 系列），或改上傳文字說明。")
        return [{"text": reading_text(Path(name).name, read_image(config, data, selected), selected["model"]),
                 "ai_vision": True}]
    pictures = slide_pictures(data)
    if not pictures:
        return []
    if not selected:
        return [{"text": f"【簡報含 {len(pictures)} 張圖片，未判讀：目前模型不支援圖片】", "ai_vision": False}]
    return [{"text": reading_text(f"第 {number} 頁圖片", read_image(config, blob, selected), selected["model"]),
             "page": number, "ai_vision": True} for number, blob in pictures]
