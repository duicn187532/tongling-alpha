"""Connector-neutral ingestion, immutable source versions, and local retrieval."""

import csv
import hashlib
import io
import re
from pathlib import Path
from typing import Protocol
import numpy as np
import faiss
from rank_bm25 import BM25Okapi
from sqlalchemy import select
from .db import Source, Chunk, Link, Project, Change, DATA, uid, log

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md", ".vtt"} | IMAGE_EXTENSIONS


class SourceProvider(Protocol):
    def health(self) -> dict: ...
    def list_objects(self) -> list: ...
    def fetch(self, identifier: str) -> bytes: ...
    def scan_changes(self, previous: dict) -> dict: ...


class InternalStorageProvider:
    def __init__(self, root=DATA / "uploads"):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)

    def health(self):
        return {"ok": self.root.is_dir()}

    def list_objects(self):
        return [p.name for p in self.root.iterdir() if p.is_file()]

    def fetch(self, identifier):
        path = (self.root / identifier).resolve()
        if not path.is_relative_to(self.root.resolve()):
            raise ValueError("Invalid object path")
        return path.read_bytes()

    def scan_changes(self, previous):
        return {
            p: hashlib.sha256(self.fetch(p)).hexdigest() for p in self.list_objects()
        }


class FileUploadProvider(InternalStorageProvider):
    def store(self, name, data):
        if Path(name).suffix.lower() not in EXTENSIONS:
            raise ValueError("不支援的檔案格式")
        if len(data) > 30 * 1024 * 1024:
            raise ValueError("檔案超過 30 MB")
        identifier = uid() + Path(name).suffix.lower()
        (self.root / identifier).write_bytes(data)
        return identifier


class LocalFolderProvider:
    def __init__(self, path):
        self.root = Path(path).resolve(strict=True)

    def health(self):
        return {"ok": self.root.is_dir()}

    def list_objects(self):
        return [
            str(p.relative_to(self.root))
            for p in self.root.rglob("*")
            if p.is_file()
            and not p.is_symlink()
            and p.suffix.lower() in EXTENSIONS - IMAGE_EXTENSIONS  # images need a vision read on upload
            and p.resolve().is_relative_to(self.root)
        ]

    def fetch(self, identifier):
        p = (self.root / identifier).resolve()
        if not p.is_relative_to(self.root):
            raise ValueError("Path outside connector root")
        if p.stat().st_size > 30 * 1024 * 1024:
            raise ValueError("檔案超過 30 MB")
        return p.read_bytes()

    def scan_changes(self, previous):
        result = {}
        for name in self.list_objects():
            p = self.root / name
            stat = p.stat()
            result[name] = {
                "mtime": stat.st_mtime_ns,
                "size": stat.st_size,
                "hash": hashlib.sha256(self.fetch(name)).hexdigest(),
            }
        return result


def parse_file(name: str, data: bytes) -> list[dict]:
    ext = Path(name).suffix.lower()
    if ext not in EXTENSIONS:
        raise ValueError("不支援的檔案格式")
    if len(data) > 30 * 1024 * 1024:
        raise ValueError("檔案超過 30 MB")
    if ext in {".docx", ".pptx", ".xlsx"}:
        import zipfile

        if not data.startswith(b"PK"):
            raise ValueError("Office 檔案內容與副檔名不符")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 150 * 1024 * 1024:
                raise ValueError("Office 解壓後大小超過上限")
    blocks = []
    if ext in IMAGE_EXTENSIONS:
        from .vision import is_image
        if not is_image(data):
            raise ValueError("圖片內容與副檔名不符")
        return []  # text comes from the vision reading (vision.image_blocks)
    if ext == ".pdf":
        import fitz

        if not data.startswith(b"%PDF"):
            raise ValueError("PDF 檔案簽章錯誤")
        with fitz.open(stream=data, filetype="pdf") as doc:
            blocks = [{"text": p.get_text(), "page": i + 1} for i, p in enumerate(doc)]
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        blocks = [{"text": p.text} for p in doc.paragraphs]
        blocks += [
            {"text": "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows)}
            for t in doc.tables
        ]
    elif ext == ".pptx":
        from pptx import Presentation

        def slide_text(slide):
            # Proposal decks keep figures in tables and context in speaker notes.
            parts = [s.text for s in slide.shapes if s.has_text_frame and s.text.strip()]
            for shape in slide.shapes:
                if shape.has_table:
                    parts.append("\n".join(" | ".join(c.text for c in row.cells) for row in shape.table.rows))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("備註：" + slide.notes_slide.notes_text_frame.text)
            return "\n".join(parts)

        blocks = [
            {"text": slide_text(slide), "page": i + 1}
            for i, slide in enumerate(Presentation(io.BytesIO(data)).slides)
        ]
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                blocks.append(
                    {
                        "text": " | ".join(
                            str(x) if x is not None else "" for x in row
                        ),
                        "heading": sheet.title,
                    }
                )
        wb.close()
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = data.decode("cp950")
            except UnicodeDecodeError as exc:
                raise ValueError("文字編碼須為 UTF-8 或 Big5") from exc
        if "\x00" in text:
            raise ValueError("不是有效文字檔案")
        if ext == ".csv":
            blocks = [
                {"text": " | ".join(row)} for row in csv.reader(io.StringIO(text))
            ]
        elif ext == ".vtt":
            for cue in re.split(r"\n\s*\n", text.replace("\r", "")):
                lines = cue.splitlines()
                timing = next(
                    (i for i, line in enumerate(lines) if "-->" in line), None
                )
                if timing is not None:
                    blocks.append(
                        {"text": " ".join(lines[timing + 1 :]), "time": lines[timing]}
                    )
        else:
            heading = ""
            for line in text.splitlines():
                if line.startswith("#"):
                    heading = line.lstrip("# ")
                if line.strip():
                    blocks.append({"text": line, "heading": heading})
    blocks = [b for b in blocks if b["text"].strip()]
    if not blocks:
        raise ValueError("沒有可解析文字；掃描 PDF 需先 OCR")
    return blocks


def mark_stale(db, project, source_id, summary):
    from .activity import cancel_project

    cancel_project(db, project.id)
    from .db import ProjectMemory

    for memory in db.scalars(
        select(ProjectMemory).where(
            ProjectMemory.project_id == project.id, ProjectMemory.status == "ACTIVE"
        )
    ):
        memory.status = "STALE"
    state = dict(project.state)
    dirty = set(state.get("dirty_sources", []))
    dirty.add(source_id)
    state["dirty_sources"] = sorted(dirty)
    state["analysis_stale"] = True
    state["draft_id"] = None
    # Review and draft reminders describe the old analysis; re-analysis is next.
    from .db import Interruption
    for row in db.scalars(select(Interruption).where(
            Interruption.project_id == project.id, Interruption.status == "OPEN",
            Interruption.kind.in_(["REQUIREMENT_REVIEW", "FINAL_PRD_REVIEW_REQUIRED"]))):
        row.status = "RESOLVED"
    added = db.get(Source, source_id)
    if added is not None and added.kind in {"REQUIREMENT_TEXT", "REQUIREMENT_DOCUMENT"}:
        # New requirement wording deserves a fresh shared-library search.
        state["retrieval_done"] = False
    project.state = state
    project.epoch += 1
    project.runtime = "READY"
    if project.current_revision_id:
        exists = db.scalar(
            select(Change).where(
                Change.project_id == project.id,
                Change.source_ref == source_id,
                Change.status.in_(["OPEN", "REVIEWED"]),
            )
        )
        if not exists:
            db.add(Change(project_id=project.id, source_ref=source_id, summary=summary))
    log(db, project.id, "EVIDENCE", summary, {"source_id": source_id})


def ingest(
    db,
    title,
    blocks,
    provider="INTERNAL_STORAGE",
    kind="USER_UPLOAD",
    project=None,
    meta=None,
):
    content = "\n".join(b["text"] for b in blocks)
    digest = hashlib.sha256(content.encode()).hexdigest()
    source = Source(
        title=title,
        content=content,
        sha256=digest,
        provider=provider,
        kind=kind,
        meta={"library_scope": "PROJECT" if project else "GLOBAL", **(meta or {})},
    )
    db.add(source)
    db.flush()
    for block in blocks:
        for pos in range(0, len(block["text"]), 680):
            db.add(
                Chunk(
                    source_id=source.id,
                    text=block["text"][pos : pos + 800],
                    meta={k: v for k, v in block.items() if k != "text"},
                )
            )
    if project:
        db.add(Link(project_id=project.id, source_id=source.id))
        mark_stale(db, project, source.id, f"加入來源：{title}")
    db.flush()
    return source


def scan_folder(db, connector):
    provider = LocalFolderProvider(connector.path)
    project = db.get(Project, connector.project_id) if connector.project_id else None
    if project and project.status != "ACTIVE":
        raise ValueError("專案已刪除，停止讀取此資料夾")
    inventory, errors, changed = {}, [], 0
    for name in provider.list_objects():
        old = connector.inventory.get(name)
        try:
            path = provider.root / name
            stat = path.stat()
            data = provider.fetch(name)
            digest = hashlib.sha256(data).hexdigest()
            if old and old["hash"] == digest:
                inventory[name] = {
                    **old,
                    "mtime": stat.st_mtime_ns,
                    "size": stat.st_size,
                }
                continue
            blocks = parse_file(name, data)
            previous = db.get(Source, old["source_id"]) if old else None
            source = ingest(
                db,
                Path(name).name,
                blocks,
                "LOCAL_FOLDER",
                kind=previous.kind
                if previous
                else "MEETING_TRANSCRIPT"
                if Path(name).suffix.lower() == ".vtt"
                else "USER_UPLOAD",
                project=project,
                meta={
                    **(previous.meta if previous else {}),
                    "connector_id": connector.id,
                    "connector_root": str(provider.root),
                    "relative_path": name,
                },
            )
            inventory[name] = {
                "mtime": stat.st_mtime_ns,
                "size": stat.st_size,
                "hash": digest,
                "source_id": source.id,
            }
            if old:
                db.get(Source, old["source_id"]).status = "SUPERSEDED"
                if project:
                    mark_stale(db, project, old["source_id"], f"來源已有新版：{name}")
                else:
                    for link in db.scalars(
                        select(Link).where(
                            Link.source_id == previous.id, Link.role == "EVIDENCE"
                        )
                    ).all():
                        affected = db.get(Project, link.project_id)
                        db.add(Link(project_id=affected.id, source_id=source.id))
                        mark_stale(db, affected, previous.id, f"來源已有新版：{name}")
                        mark_stale(db, affected, source.id, f"加入來源新版：{name}")
            changed += 1
        except Exception as exc:
            errors.append({"file": name, "error": str(exc)})
            if old:
                inventory[name] = old
    removed = set(connector.inventory) - set(inventory)
    for name in removed:
        source = db.get(Source, connector.inventory[name]["source_id"])
        source.status = "REMOVED"
        if project:
            mark_stale(db, project, source.id, f"資料夾移除來源：{name}")
        else:
            for link in db.scalars(
                select(Link).where(Link.source_id == source.id, Link.role == "EVIDENCE")
            ).all():
                mark_stale(
                    db,
                    db.get(Project, link.project_id),
                    source.id,
                    f"資料夾移除來源：{name}",
                )
    connector.inventory, connector.errors = inventory, errors
    return {"changed": changed, "removed": len(removed), "errors": errors}


def tokens(text):
    chars = re.findall(r"[\u3400-\u9fff]", text)
    return (
        re.findall(r"[a-z0-9_]+", text.lower())
        + chars
        + ["".join(chars[i : i + 2]) for i in range(len(chars) - 1)]
    )


def vectors(texts, model_path=""):
    if model_path:
        from sentence_transformers import SentenceTransformer

        model = SentenceTransformer(model_path, local_files_only=True)
        return np.asarray(
            model.encode(texts, normalize_embeddings=True), dtype="float32"
        )
    # Offline lexical feature embeddings: no implicit network/model downloads.
    result = np.zeros((len(texts), 512), dtype="float32")
    for i, text in enumerate(texts):
        for token in tokens(text):
            index = int(hashlib.sha256(token.encode()).hexdigest()[:8], 16) % 512
            result[i, index] += 1
    faiss.normalize_L2(result)
    return result


def search(db, query, limit=8, model_path="", eligible_ids=None, strategy="hybrid"):
    if eligible_ids is not None and not eligible_ids:
        return []
    statement = (select(Chunk, Source).join(Source, Chunk.source_id == Source.id)
                 .where(Source.status == "ACTIVE"))
    if eligible_ids is not None:
        statement = statement.where(Chunk.source_id.in_(eligible_ids))
    rows = db.execute(statement).all()
    rows = [
        (c, s)
        for c, s in rows
        if (
            s.id in eligible_ids
            if eligible_ids is not None
            else s.meta.get("library_scope") == "GLOBAL"
        )
        and s.kind != "MEETING_TRANSCRIPT"
    ]
    if not rows or not tokens(query):
        return []
    corpus = [tokens(c.text) for c, _ in rows]
    bm = BM25Okapi(corpus).get_scores(tokens(query))
    matrix = vectors([c.text for c, _ in rows] + [query], model_path)
    index = faiss.IndexFlatIP(matrix.shape[1])
    index.add(matrix[:-1])
    scores, indices = index.search(matrix[-1:], min(20, len(rows)))
    rank = {}
    for position, i in enumerate(np.argsort(-bm)[:20]):
        if strategy != "vector" and set(tokens(query)) & set(corpus[i]):
            rank[int(i)] = 1 / (60 + position + 1)
    for position, (i, score) in enumerate(zip(indices[0], scores[0])):
        if strategy != "bm25" and score > 0.05:
            rank[int(i)] = rank.get(int(i), 0) + 1 / (60 + position + 1)
    results = []
    for i, score in sorted(rank.items(), key=lambda x: -x[1])[:limit]:
        chunk, source = rows[i]
        results.append(
            {
                "source_id": source.id,
                "chunk_id": chunk.id,
                "title": source.title,
                "text": chunk.text,
                "meta": {**source.meta, **chunk.meta},
                "kind": source.kind,
                "score": round(score, 5),
                "reason": "FAISS 向量 + BM25 排名融合（非可信度機率）",
            }
        )
    return results
