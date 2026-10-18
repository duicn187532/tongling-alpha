"""Loopback-only file server implementing TongLing's knowledge API contract.

This process owns its own folder and index. TongLing sees it only through HTTP.
"""

import hashlib
import mimetypes
import os
import re
import threading
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel, Field


ROOT = Path(__file__).resolve().parents[1]
FILES = Path(os.getenv("COMPANY_FILES_DIR", ROOT / "company_files")).resolve()
SAMPLES = ROOT / "sample_data" / "company_file_server"
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".jsonl", ".yaml", ".yml", ".xml", ".sql", ".log", ".cyjs", ".mmd", ".mermaid", ".dot", ".gml"}
GRAPH_EXTENSIONS = {".json", ".cyjs", ".mmd", ".mermaid", ".dot", ".gml", ".graphml", ".gexf", ".drawio", ".svg", ".vsdx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
EXTENSIONS = TEXT_EXTENSIONS | GRAPH_EXTENSIONS | IMAGE_EXTENSIONS | {".pdf", ".docx", ".pptx", ".xlsx", ".odt", ".ods", ".odp"}
MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_TEXT_CHARS = 1_000_000


def _terms(value):
    result = set(re.findall(r"[a-z0-9]+", value.lower()))
    for run in re.findall(r"[\u3400-\u9fff]+", value):
        result.update(run[i:i + 2] for i in range(max(1, len(run) - 1)))
        if len(run) == 1:
            result.add(run)
    return result


def _xml(data):
    if len(data) > 5_000_000 or b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ValueError("XML file is too large or contains unsupported declarations")
    return ElementTree.fromstring(data)


def _text(path):
    extension = path.suffix.lower()
    if extension in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8-sig", errors="replace")
    if extension in IMAGE_EXTENSIONS:
        from PIL import Image
        with Image.open(path) as picture:
            width, height = picture.size
            details = [f"圖片尺寸 {width}×{height}", f"格式 {picture.format}"]
            for value in picture.getexif().values():
                if isinstance(value, str) and len(value) < 300:
                    details.append(value)
            return "；".join(details)
    if extension in {".graphml", ".gexf", ".drawio", ".svg"}:
        root = _xml(path.read_bytes())
        parts = [value for node in root.iter() for value in
                 [node.text, *(str(item) for item in node.attrib.values())]
                 if value and value.strip()]
        return "\n".join(parts)
    if extension in {".odt", ".ods", ".odp", ".vsdx"}:
        with zipfile.ZipFile(path) as package:
            names = [name for name in package.namelist()
                     if name == "content.xml" or (extension == ".vsdx" and name.startswith("visio/pages/page") and name.endswith(".xml"))]
            parts = []
            for name in names[:100]:
                if package.getinfo(name).file_size > MAX_TEXT_CHARS * 5:
                    continue
                parts.extend(t.strip() for t in _xml(package.read(name)).itertext() if t.strip())
            return "\n".join(parts)
    if extension == ".pdf":
        import pymupdf
        with pymupdf.open(path) as document:
            return "\n".join(page.get_text() for page in document)
    if extension == ".docx":
        from docx import Document
        document = Document(path)
        return "\n".join([p.text for p in document.paragraphs] +
                         [" | ".join(cell.text for row in table.rows for cell in row.cells)
                          for table in document.tables])
    if extension == ".pptx":
        from pptx import Presentation
        return "\n".join(shape.text for slide in Presentation(path).slides
                         for shape in slide.shapes if shape.has_text_frame)
    if extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return "\n".join(" | ".join(str(value) for value in row if value is not None)
                             for sheet in workbook for row in sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    return ""


class FileIndex:
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.docs = {}
        self.errors = {}
        self.stamps = {}
        self.last_scan = 0.0
        self.lock = threading.RLock()

    def refresh(self, force=False):
        with self.lock:
            if not force and time.monotonic() - self.last_scan < 3:
                return
            seen = set()
            errors = {}
            for file in self.root.rglob("*"):
                if not file.is_file() or file.is_symlink():
                    continue
                try:
                    if not file.resolve().is_relative_to(self.root):
                        continue
                    relative = file.relative_to(self.root).as_posix()
                    seen.add(relative)
                    info = file.stat()
                    stamp = (info.st_mtime_ns, info.st_size)
                    if stamp == self.stamps.get(relative):
                        continue
                    if info.st_size > MAX_FILE_BYTES:
                        raise ValueError("File exceeds 30 MB limit")
                    content = _text(file)[:MAX_TEXT_CHARS]
                    identifier = hashlib.sha256(relative.encode()).hexdigest()[:24]
                    digest = hashlib.sha256(file.read_bytes()).hexdigest()
                    category = ("IMAGE" if file.suffix.lower() in IMAGE_EXTENSIONS else
                                "GRAPH" if file.suffix.lower() in GRAPH_EXTENSIONS else
                                "DOCUMENT" if file.suffix.lower() in EXTENSIONS else "FILE")
                    self.docs[relative] = {
                        "id": identifier, "title": file.name, "relative_path": relative,
                        "content": content, "content_hash": digest,
                        "version": digest[:12], "terms": _terms(content + " " + file.name),
                        "category": category, "file_type": file.suffix.lower().lstrip("."),
                        "media_type": mimetypes.guess_type(file.name)[0] or "application/octet-stream",
                        "content_mode": "METADATA_ONLY" if category in {"IMAGE", "FILE"} else "EXTRACTED_TEXT",
                    }
                    self.stamps[relative] = stamp
                except Exception as exc:
                    relative = file.relative_to(self.root).as_posix()
                    self.docs.pop(relative, None)
                    self.stamps.pop(relative, None)
                    errors[relative] = str(exc)[:180]
            for relative in set(self.docs) - seen:
                self.docs.pop(relative, None)
                self.stamps.pop(relative, None)
            self.errors = errors
            self.last_scan = time.monotonic()

    def list(self):
        self.refresh()
        with self.lock:
            return sorted(self.docs.values(), key=lambda d: d["relative_path"])

    def by_id(self, identifier):
        return next((doc for doc in self.list() if doc["id"] == identifier), None)


def _seed_samples():
    FILES.mkdir(parents=True, exist_ok=True)
    if any(FILES.iterdir()) or not SAMPLES.is_dir():
        return
    for sample in SAMPLES.iterdir():
        if sample.is_file() and sample.name != "README.md":
            (FILES / sample.name).write_bytes(sample.read_bytes())


_seed_samples()
index = FileIndex(FILES)
app = FastAPI(title="Company File Server Simulator", docs_url=None, redoc_url=None, openapi_url=None)


@app.middleware("http")
async def local_only(request: Request, call_next):
    if request.url.hostname not in {"localhost", "127.0.0.1", "testserver"}:
        return JSONResponse({"detail": "Local file server only"}, status_code=400)
    return await call_next(request)


class SearchRequest(BaseModel):
    query: str = Field(min_length=1, max_length=1000)
    limit: int = Field(default=5, ge=1, le=20)


def _summary(doc):
    return {"id": doc["id"], "title": doc["title"], "relative_path": doc["relative_path"],
            "version": doc["version"], "characters": len(doc["content"]),
            "uri": "company-file://" + doc["relative_path"], "category": doc["category"],
            "file_type": doc["file_type"], "media_type": doc["media_type"],
            "content_mode": doc["content_mode"]}


def _metadata(doc):
    return {"version": doc["version"], "status": "ACTIVE",
            "canonical_uri": "company-file://" + doc["relative_path"],
            "content_hash": doc["content_hash"], "authority_class": "UNVERIFIED_INTERNAL",
            "verified_authority": False, "category": doc["category"],
            "file_type": doc["file_type"], "content_mode": doc["content_mode"]}


@app.get("/health")
def health():
    index.refresh()
    return {"status": "ok", "service": "company-file-server", "documents": len(index.docs),
            "errors": len(index.errors), "folder": str(FILES)}


@app.get("/capabilities")
def capabilities():
    return {"capabilities": ["SEARCH", "KEYWORD", "GRAPH", "FETCH_SNIPPET", "METADATA", "VERSION", "CHANGE_SCAN"]}


@app.get("/documents")
def documents():
    return {"documents": [_summary(doc) for doc in index.list()], "errors": index.errors}


@app.post("/rescan")
def rescan():
    index.refresh(force=True)
    return {"documents": len(index.docs), "errors": index.errors}


def _rank(body: SearchRequest, category=None):
    query_terms = _terms(body.query)
    candidates = []
    for doc in index.list():
        if category and doc["category"] != category:
            continue
        matched = query_terms & doc["terms"]
        if not matched:
            continue
        title_match = query_terms & _terms(doc["title"])
        score = min(1.0, 0.8 * len(matched) / max(1, len(query_terms)) +
                    0.2 * len(title_match) / max(1, len(query_terms)))
        content = doc["content"]
        lower = content.lower()
        offset = min((lower.find(term.lower()) for term in matched if lower.find(term.lower()) >= 0), default=0)
        snippet = content[max(0, offset - 80):offset + 360]
        if doc["category"] in {"IMAGE", "FILE"}:
            snippet = ("圖片" if doc["category"] == "IMAGE" else "附件") + f"原檔待人工檢視；目前只可搜尋檔名與中繼資料。{content}"
        candidates.append({"id": doc["id"], "title": doc["title"], "snippet": snippet,
                           "score": round(score, 4), "uri": "company-file://" + doc["relative_path"],
                           "metadata": _metadata(doc)})
    candidates.sort(key=lambda item: (-item["score"], item["title"]))
    return candidates[:body.limit]


@app.post("/search")
def search(body: SearchRequest):
    return {"results": _rank(body)}


@app.post("/graph/search")
def graph_search(body: SearchRequest):
    return {"results": _rank(body, "GRAPH")}


@app.get("/evidence/{identifier}")
def evidence(identifier: str):
    doc = index.by_id(identifier)
    if not doc:
        raise HTTPException(404, "Document not found")
    return {**_summary(doc), "content": doc["content"], "metadata": _metadata(doc)}


@app.get("/evidence/{identifier}/metadata")
def evidence_metadata(identifier: str):
    doc = index.by_id(identifier)
    if not doc:
        raise HTTPException(404, "Document not found")
    return {**_summary(doc), "metadata": _metadata(doc)}


@app.get("/files/{identifier}")
def original_file(identifier: str):
    doc = index.by_id(identifier)
    if not doc:
        raise HTTPException(404, "Document not found")
    original = index.root / doc["relative_path"]
    path = original.resolve()
    if original.is_symlink() or not path.is_relative_to(index.root) or not path.is_file():
        raise HTTPException(404, "Document not found")
    return FileResponse(path, media_type=doc["media_type"], filename=doc["title"], content_disposition_type="attachment")


@app.get("/preview/{identifier}")
def preview_image(identifier: str):
    doc = index.by_id(identifier)
    if not doc or doc["file_type"] not in {"png", "jpg", "jpeg", "webp", "gif"}:
        raise HTTPException(404, "Preview unavailable")
    original = index.root / doc["relative_path"]
    path = original.resolve()
    if original.is_symlink() or not path.is_relative_to(index.root) or not path.is_file():
        raise HTTPException(404, "Preview unavailable")
    return FileResponse(path, media_type=doc["media_type"], headers={"Content-Security-Policy": "default-src 'none'", "X-Content-Type-Options": "nosniff"})
